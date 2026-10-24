"""Word exports on a Nojoin letterhead.

Mirrors the PDF layout: a letterhead band on the first page, a slim running
header after it, and a footer with the page count on every page. Charts are
the same scenes the PDF draws, painted to PNG because Word cannot take vector
drawings from here. Geist is embedded so the document reads in the brand type
on machines that do not have it.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Sequence
from dataclasses import fields, is_dataclass, replace
from io import BytesIO
from typing import Any, cast

from docx import Document as new_document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.shared import Mm, Pt, RGBColor, Twips
from docx.text.run import Run

from . import docx_parts, theme
from .charts import layout
from .docx_fonts import embed_geist
from .formatting import format_transcript_time
from .model import (
    Block,
    Callout,
    ChartBlock,
    CodeBlock,
    ExportDocument,
    Heading,
    Inline,
    ListBlock,
    Paragraph,
    Rule,
    Section,
    Span,
    StatRow,
    Table,
    TranscriptLine,
    inline_text,
)
from .paint import paint_png
from .render_pdf import column_widths

PAGE_MARGIN = Mm(18)
TEXT_WIDTH_PT = (210 - 36) / 25.4 * 72
_SAFE_LINK_PREFIXES = ("http://", "https://", "mailto:")
_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")
_NUMBER_STYLES = ("List Number", "List Number 2", "List Number 3")
_MAX_HEADING_LEVEL = 9
# Geist Mono is not embedded (see docx_fonts.py), so code falls back to a font
# every Word install has rather than to whatever the reader's machine picks.
_CODE_FONT = "Courier New"
_TEMPLATE_HEADER_STOPS = (Twips(4680), Twips(9360))
# Everything outside XML 1.0's Char production: the C0 controls other than tab,
# newline and carriage return, lone surrogates, U+FFFE and U+FFFF. lxml refuses
# a string holding any of them, so one in a transcript line or a speaker's name
# would fail the whole export (#685).
_XML_INVALID = re.compile(r"[^\t\n\r\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")


def xml_safe(text: str) -> str:
    """The text without the characters a Word document cannot hold."""
    return _XML_INVALID.sub("", text)


def _xml_safe_fields(value: Any) -> Any:
    """The value with every string in it made safe, at any depth."""
    if isinstance(value, str):
        return xml_safe(value)
    if isinstance(value, tuple):
        return tuple(_xml_safe_fields(item) for item in value)
    if is_dataclass(value) and not isinstance(value, type):
        return replace(
            value,
            **{
                item.name: _xml_safe_fields(getattr(value, item.name))
                for item in fields(value)
            },
        )
    return value


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


# -- runs ---------------------------------------------------------------------------


def _style_run(run: Any, span: Span, *, size: float | None, color: str | None) -> None:
    run.bold = span.bold or None
    run.italic = span.italic or None
    if span.strike:
        run.font.strike = True
    if span.code:
        run.font.name = _CODE_FONT
        run.font.size = Pt(8.5)
    elif size:
        run.font.size = Pt(size)
    if span.color or color:
        run.font.color.rgb = _rgb(span.color or color or theme.FOREGROUND)


def add_spans(
    paragraph: Any,
    spans: Inline,
    *,
    size: float | None = None,
    color: str | None = None,
) -> None:
    for span in spans:
        if span.text == "\n":
            paragraph.add_run().add_break()
            continue
        if span.link and span.link.lower().startswith(_SAFE_LINK_PREFIXES):
            hyperlink = docx_parts.add_hyperlink(paragraph, span.link)
            if hyperlink is not None:
                run = paragraph.add_run(span.text)
                _style_run(run, span, size=size, color=color)
                run.underline = True
                run.font.color.rgb = _rgb(theme.BRAND)
                hyperlink.append(run._r)
                continue
        run = paragraph.add_run(span.text)
        _style_run(run, span, size=size, color=color)


# -- document setup -----------------------------------------------------------------


def _configure_styles(doc: Any) -> None:
    normal = doc.styles["Normal"]
    docx_parts.set_font_name(normal.element, theme.FONT_REGULAR)
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(theme.FOREGROUND)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    sizes = {1: 16, 2: 13, 3: 11, 4: 10}
    for level in range(1, 10):
        style = doc.styles[f"Heading {level}"]
        docx_parts.set_font_name(style.element, theme.FONT_REGULAR)
        style.font.bold = True
        style.font.italic = False
        style.font.size = Pt(sizes.get(level, 10))
        style.font.color.rgb = _rgb(theme.FOREGROUND if level < 4 else theme.MUTED)
        docx_parts.clear_theme_color(style.element.get_or_add_rPr())
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(4)
    for name in (*_BULLET_STYLES, *_NUMBER_STYLES, "Quote"):
        try:
            docx_parts.set_font_name(doc.styles[name].element, theme.FONT_REGULAR)
        except KeyError:
            continue


def _configure_page(section: Any) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = PAGE_MARGIN
    section.top_margin = Mm(30)
    section.bottom_margin = Mm(22)
    section.header_distance = Mm(11)
    section.footer_distance = Mm(9)
    section.different_first_page_header_footer = True


def _right_tab(paragraph: Any) -> None:
    stops = paragraph.paragraph_format.tab_stops
    # The template's Header and Footer styles set centre and right stops for
    # Letter paper; left in place, the first tab lands mid-page.
    for position in _TEMPLATE_HEADER_STOPS:
        stops.add_tab_stop(position, WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Pt(TEXT_WIDTH_PT), WD_TAB_ALIGNMENT.RIGHT)


def _first_page_header(section: Any, document: ExportDocument) -> None:
    accent = theme.accent(document.branded)
    paragraph = section.first_page_header.paragraphs[0]
    _right_tab(paragraph)
    if document.branded:
        paragraph.add_run().add_picture(str(theme.LOGO_PATH), width=Pt(20))
        name = paragraph.add_run("  Nojoin")
        name.bold = True
        name.font.size = Pt(15)
    label = paragraph.add_run("\tMEETING RECORD")
    label.bold = True
    label.font.size = Pt(7)
    label.font.color.rgb = _rgb(theme.HELPER)
    dated = section.first_page_header.add_paragraph()
    dated.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = dated.add_run(document.dated)
    run.font.size = Pt(9.5)
    run.font.color.rgb = _rgb(theme.MUTED)
    docx_parts.paragraph_border(dated, "bottom", size=12, color=accent, space=6)


def _later_header(section: Any, document: ExportDocument) -> None:
    paragraph = section.header.paragraphs[0]
    _right_tab(paragraph)
    if document.branded:
        paragraph.add_run().add_picture(str(theme.LOGO_PATH), width=Pt(11))
        paragraph.add_run("  ")
    title = paragraph.add_run(document.title)
    title.font.size = Pt(8)
    title.font.color.rgb = _rgb(theme.HELPER)
    dated = paragraph.add_run(f"\t{document.dated}")
    dated.font.size = Pt(8)
    dated.font.color.rgb = _rgb(theme.HELPER)
    docx_parts.paragraph_border(
        paragraph, "bottom", size=4, color=theme.HAIRLINE, space=4
    )


def _footer(footer: Any, document: ExportDocument) -> None:
    paragraph = footer.paragraphs[0]
    _right_tab(paragraph)
    docx_parts.paragraph_border(paragraph, "top", size=4, color=theme.HAIRLINE, space=4)
    if document.branded:
        add_spans(
            paragraph,
            (
                Span(f"{theme.GENERATED_BY} · ", color=theme.HELPER),
                Span(theme.SITE_LABEL, link=theme.SITE_URL),
            ),
            size=7.5,
        )
    for text, field in (("\tPage ", "PAGE"), (" of ", "NUMPAGES")):
        label = paragraph.add_run(text)
        number = Run(docx_parts.add_field(paragraph, field, "1"), paragraph)
        for run in (label, number):
            run.font.size = Pt(7.5)
            run.font.color.rgb = _rgb(theme.HELPER)


# -- blocks ---------------------------------------------------------------------------


class _BodyWriter:
    """Writes blocks into a container: the document body or a table cell."""

    def __init__(self, doc: Any, accent: str) -> None:
        self.doc = doc
        self.accent = accent
        self._renderers: dict[type, Callable[[Any, Any], None]] = {
            Heading: self._heading,
            Paragraph: self._paragraph,
            ListBlock: self._list,
            Table: self._table,
            Rule: self._rule,
            CodeBlock: self._code,
            StatRow: self._stats,
            Callout: self._callout,
            ChartBlock: self._chart,
            TranscriptLine: self._transcript,
        }

    def write(self, container: Any, blocks: Sequence[Block]) -> None:
        for block in blocks:
            self._renderers[type(block)](container, block)

    def section_title(self, title: str) -> None:
        paragraph = self.doc.add_heading(title, level=1)
        docx_parts.paragraph_border(
            paragraph, "bottom", size=12, color=self.accent, space=3
        )

    def _heading(self, container: Any, block: Heading) -> None:
        level = min(block.level + 1, _MAX_HEADING_LEVEL)
        paragraph = container.add_paragraph(style=f"Heading {level}")
        add_spans(paragraph, block.text)

    def _paragraph(self, container: Any, block: Paragraph) -> None:
        tone = {
            "muted": (9.0, theme.HELPER),
            "small": (8.0, theme.FAINT),
            "quote": (None, theme.MUTED),
        }.get(block.tone, (None, None))
        paragraph = container.add_paragraph(
            style="Quote" if block.tone == "quote" else None
        )
        add_spans(paragraph, block.text, size=tone[0], color=tone[1])

    def _list(self, container: Any, block: ListBlock) -> None:
        numbering: dict[int, int | None] = {}
        for item in block.items:
            styles = _NUMBER_STYLES if item.ordered else _BULLET_STYLES
            style = styles[min(item.depth, len(styles) - 1)]
            paragraph = container.add_paragraph(style=style)
            if item.ordered:
                if item.depth not in numbering:
                    numbering[item.depth] = docx_parts.restart_numbering(
                        self.doc, style
                    )
                num_id = numbering[item.depth]
                if num_id is not None:
                    docx_parts.apply_numbering(paragraph, num_id)
            add_spans(paragraph, item.text)

    def _table(self, container: Any, block: Table) -> None:
        has_header = any(inline_text(cell) for cell in block.header)
        rows = ([block.header] if has_header else []) + list(block.rows)
        table = container.add_table(rows=len(rows), cols=block.column_count)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        docx_parts.set_table_borders(
            table,
            {"bottom": (4, theme.HAIRLINE), "insideH": (4, theme.HAIRLINE)},
        )
        docx_parts.set_table_cell_margins(table, top=70, bottom=70, left=100, right=100)
        widths = column_widths(block, TEXT_WIDTH_PT)
        docx_parts.set_grid_widths(table, widths)
        offset = 1 if has_header else 0
        for row_index, row in enumerate(rows):
            header = has_header and row_index == 0
            for column in range(block.column_count):
                cell = table.cell(row_index, column)
                cell.width = Pt(widths[column])
                self._table_cell(
                    cell,
                    row[column] if column < len(row) else (),
                    header,
                    column in block.numeric,
                )
            if header:
                docx_parts.repeat_as_header(table.rows[0])
            marker = row_index - offset
            if 0 <= marker < len(block.row_colors) and block.row_colors[marker]:
                docx_parts.set_cell_borders(
                    table.cell(row_index, 0),
                    {"left": (18, block.row_colors[marker] or "")},
                )
        container.add_paragraph().paragraph_format.space_after = Pt(2)

    def _table_cell(
        self, cell: Any, spans: Inline, header: bool, numeric: bool
    ) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        if numeric:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if header:
            docx_parts.shade_cell(cell, theme.TABLE_HEADER_FILL)
            docx_parts.set_cell_borders(cell, {"bottom": (10, self.accent)})
            add_spans(
                paragraph,
                tuple(Span(s.text, bold=True) for s in spans),
                size=8.5,
                color=theme.MUTED,
            )
        else:
            add_spans(paragraph, spans, size=9)

    def _rule(self, container: Any, block: Rule) -> None:
        paragraph = container.add_paragraph()
        docx_parts.paragraph_border(
            paragraph, "bottom", size=4, color=theme.RULE, space=1
        )

    def _code(self, container: Any, block: CodeBlock) -> None:
        paragraph = container.add_paragraph()
        add_spans(paragraph, (Span(block.text, code=True),), color=theme.MUTED)

    def _stats(self, container: Any, block: StatRow) -> None:
        table = container.add_table(rows=1, cols=max(len(block.stats), 1))
        table.style = "Table Grid"
        docx_parts.set_table_borders(
            table,
            {
                side: (4, theme.HAIRLINE)
                for side in ("top", "left", "bottom", "right", "insideV")
            },
        )
        docx_parts.set_table_cell_margins(
            table, top=110, bottom=110, left=160, right=120
        )
        for cell, stat in zip(table.rows[0].cells, block.stats):
            docx_parts.shade_cell(cell, theme.TABLE_HEADER_FILL)
            label = cell.paragraphs[0]
            label.paragraph_format.space_after = Pt(1)
            add_spans(label, (Span(stat.label),), size=7.5, color=theme.HELPER)
            value = cell.add_paragraph()
            value.paragraph_format.space_after = Pt(0)
            add_spans(value, (Span(stat.value, bold=True),), size=15)
        container.add_paragraph().paragraph_format.space_after = Pt(2)

    def _callout(self, container: Any, block: Callout) -> None:
        fill, border, text = theme.CALLOUT_COLORS[block.tone]
        table = container.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        docx_parts.set_table_borders(table, {})
        docx_parts.set_table_cell_margins(
            table, top=110, bottom=60, left=180, right=160
        )
        cell = table.cell(0, 0)
        docx_parts.shade_cell(cell, fill)
        docx_parts.set_cell_borders(cell, {"left": (18, border)})
        docx_parts.keep_row_together(table.rows[0])
        first = cell.paragraphs[0]
        if block.title:
            add_spans(first, (Span(block.title, bold=True),), color=text)
        else:
            first._p.getparent().remove(first._p)
        self.write(cell, block.body)
        container.add_paragraph().paragraph_format.space_after = Pt(2)

    def _chart(self, container: Any, block: ChartBlock) -> None:
        scene = layout(block.chart, TEXT_WIDTH_PT)
        paragraph = container.add_paragraph()
        paragraph.add_run().add_picture(
            BytesIO(paint_png(scene)), width=Pt(scene.width)
        )
        paragraph.paragraph_format.space_after = Pt(10)

    def _transcript(self, container: Any, block: TranscriptLine) -> None:
        paragraph = container.add_paragraph()
        layout_format = paragraph.paragraph_format
        # A hanging indent keeps wrapped lines clear of the timestamp.
        layout_format.left_indent = Pt(44)
        layout_format.first_line_indent = Pt(-44)
        layout_format.space_after = Pt(3)
        layout_format.tab_stops.add_tab_stop(Pt(44))
        stamp = paragraph.add_run(format_transcript_time(block.start_seconds) + "\t")
        stamp.font.size = Pt(8)
        stamp.font.color.rgb = _rgb(theme.FAINT)
        name = paragraph.add_run(f"{block.speaker}: ")
        name.bold = True
        name.font.size = Pt(9.5)
        if block.color:
            name.font.color.rgb = _rgb(theme.readable_on_white(block.color))
        text = paragraph.add_run(block.text)
        text.font.size = Pt(9.5)


def _title_block(doc: Any, document: ExportDocument) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    add_spans(title, (Span(document.title, bold=True),), size=21)
    meta = {item.label: item.value for item in document.meta}
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(1)
    first = "  ·  ".join(
        value for key in ("Date", "Duration") if (value := meta.get(key))
    )
    add_spans(line, (Span(first),), size=9.5, color=theme.MUTED)
    if meta.get("Speakers"):
        speakers = doc.add_paragraph()
        add_spans(
            speakers,
            (Span("Speakers  ", bold=True), Span(meta["Speakers"])),
            size=9.5,
            color=theme.MUTED,
        )
        line = speakers
    docx_parts.paragraph_border(line, "bottom", size=4, color=theme.HAIRLINE, space=8)


def _set_properties(doc: Any, document: ExportDocument) -> None:
    properties = doc.core_properties
    properties.title = document.title
    properties.author = "Nojoin" if document.branded else ""
    properties.last_modified_by = properties.author
    # python-docx's template says it was generated by python-docx.
    properties.comments = ""
    properties.keywords = ""


def _write_sections(doc: Any, writer: _BodyWriter, sections: Sequence[Section]) -> None:
    for index, section in enumerate(sections):
        if index:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        writer.section_title(section.title)
        writer.write(doc, section.blocks)


def render_docx(document: ExportDocument) -> bytes:
    # Cleaned once here rather than where each string reaches python-docx, so
    # no write path (core properties, headers, transcript runs) can miss it.
    document = cast(ExportDocument, _xml_safe_fields(document))
    doc = new_document()
    accent = theme.accent(document.branded)
    _configure_styles(doc)
    section = doc.sections[0]
    _configure_page(section)
    section.start_type = WD_SECTION.NEW_PAGE
    _first_page_header(section, document)
    _later_header(section, document)
    _footer(section.first_page_footer, document)
    _footer(section.footer, document)
    _set_properties(doc, document)

    _title_block(doc, document)
    _write_sections(doc, _BodyWriter(doc, accent), document.sections)
    embed_geist(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
