"""Meeting notes, which are stored as Markdown, in a Word export.

The exporter these cover replaced a line-by-line parser that emitted a Markdown
table into Word as paragraphs of pipes and hyphens (issue #136). Notes now go
through the shared export pipeline (issue #419), so these render a whole
document and read the notes back out of it.
"""

from io import BytesIO

from docx import Document

from backend.services.exports import build_document, render
from backend.tests.export_fixtures import meeting_parts

DECISION_TABLE = """## Key Decisions

| ID | Decision | Rationale |
| --- | --- | --- |
| DEC-001 | Use **PostgreSQL** | Architecture alignment |
| DEC-002 | Keep Memgraph | Graph projection |
"""


def _render(markdown_text: str):
    document = build_document(meeting_parts(notes=markdown_text), ["notes"])
    return Document(BytesIO(render(document, "docx")))


def _notes_paragraphs(document) -> list:
    """Body paragraphs after the "Meeting notes" section heading."""
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text == "Meeting notes")
    return paragraphs[start + 1 :]


def test_markdown_table_becomes_a_native_word_table() -> None:
    document = _render(DECISION_TABLE)

    assert len(document.tables) == 1
    table = document.tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 3)
    assert [cell.text for cell in table.rows[0].cells] == [
        "ID",
        "Decision",
        "Rationale",
    ]
    assert [cell.text for cell in table.rows[2].cells] == [
        "DEC-002",
        "Keep Memgraph",
        "Graph projection",
    ]


def test_table_pipes_never_reach_the_document_body() -> None:
    # The exact failure reported: rows and delimiters exported as paragraphs.
    document = _render(DECISION_TABLE)

    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "|" not in body_text
    assert "---" not in body_text


def test_table_has_a_bold_header_row_that_repeats_across_pages() -> None:
    document = _render(DECISION_TABLE)
    table = document.tables[0]

    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    # tblHeader makes Word repeat the header row across a page break.
    assert "tblHeader" in table.rows[0]._tr.xml
    assert "tblHeader" not in table.rows[1]._tr.xml


def test_inline_formatting_survives_inside_a_cell() -> None:
    document = _render(DECISION_TABLE)
    cell = document.tables[0].rows[1].cells[1]

    assert cell.text == "Use PostgreSQL"
    assert any(run.bold for run in cell.paragraphs[0].runs)


def test_line_break_in_a_cell_is_rendered_as_a_break() -> None:
    # Markdown has no in-cell newline, so both the editor and the notes prompt
    # use <br>. A real newline there would split the row.
    document = _render(
        "| Action | Owner |\n| --- | --- |\n| Prep diagram<br>and circulate | Peter |"
    )

    assert document.tables[0].rows[1].cells[0].text == "Prep diagram\nand circulate"


def test_a_six_column_table_keeps_every_column() -> None:
    document = _render(
        "| ID | Action | Owner | Contributors | Target date | Jira |\n"
        "| --- | --- | --- | --- | --- | --- |\n"
        "| ACT-001 | Prepare the diagram | Peter | Anna, Jon | 2026-08-01 | ENG-1421 |"
    )

    table = document.tables[0]
    assert len(table.columns) == 6
    assert table.rows[1].cells[5].text == "ENG-1421"


def test_ragged_rows_are_padded_rather_than_dropped() -> None:
    document = _render("| A | B | C |\n| --- | --- | --- |\n| only one |")

    table = document.tables[0]
    assert len(table.columns) == 3
    assert [cell.text for cell in table.rows[1].cells] == ["only one", "", ""]


def test_headings_lists_and_marks_render_as_word_elements() -> None:
    document = _render(
        "## Detailed Notes\n\n"
        "### Topic\n\n"
        "- Top level\n"
        "  - Nested\n\n"
        "1. First\n\n"
        "Body with **bold**, *italic* and ~~struck~~ text.\n"
    )

    styles = [
        (paragraph.style.name, paragraph.text)
        for paragraph in _notes_paragraphs(document)
        if paragraph.text.strip()
    ]

    # The notes' shallowest heading sits one level below the section title.
    assert ("Heading 2", "Detailed Notes") in styles
    assert ("Heading 3", "Topic") in styles
    assert ("List Bullet", "Top level") in styles
    assert ("List Bullet 2", "Nested") in styles
    assert ("List Number", "First") in styles

    body = next(p for p in document.paragraphs if p.text.startswith("Body with"))
    assert any(run.bold for run in body.runs)
    assert any(run.italic for run in body.runs)
    assert any(run.font.strike for run in body.runs)
    # Strikethrough must be parsed, not passed through as punctuation.
    assert "~~" not in body.text


def test_each_numbered_list_counts_from_one() -> None:
    # Word shares one numbering instance per list style unless told otherwise,
    # so a second list would carry on from the first.
    document = _render("1. One\n2. Two\n\nBetween.\n\n1. Again\n")

    numbered = [p for p in _notes_paragraphs(document) if p.style.name == "List Number"]
    num_ids = {p._p.pPr.numPr.numId.val for p in numbered}
    assert len(numbered) == 3
    assert len(num_ids) == 2


def test_links_become_real_hyperlinks() -> None:
    document = _render("See [the plan](https://example.com/plan) for detail.")

    paragraph = next(p for p in document.paragraphs if p.text.startswith("See"))
    assert "w:hyperlink" in paragraph._p.xml
    assert "https://example.com/plan" not in paragraph.text


def test_only_web_and_mail_links_become_hyperlinks() -> None:
    # markdown-it already refuses javascript: and data: URLs; the renderer's
    # own allow-list also keeps any other scheme out of the document.
    document = _render("Fetch [the file](ftp://example.com/file) here.")

    paragraph = next(p for p in document.paragraphs if p.text.startswith("Fetch"))
    assert "w:hyperlink" not in paragraph._p.xml
    assert paragraph.text == "Fetch the file here."


def test_empty_notes_say_so_rather_than_rendering_nothing() -> None:
    document = _render("")

    text = "\n".join(cell.text for table in document.tables for cell in table._cells)
    assert "No meeting notes have been generated" in text
