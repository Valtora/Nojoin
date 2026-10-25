"""Orchestrator behaviour: vision fan-out, downgrade, resume, and staleness.

These exercise the decision logic around the parse rather than the parsers
themselves (covered in test_document_parsing.py), using stub backends so no
provider is contacted.
"""

from __future__ import annotations

import io

import pytest

from backend.models.document_page import PageParseMode
from backend.processing.documents.types import PageSource
from backend.processing.llm_backends.base import (
    LLMBackend,
    is_vision_unsupported_error,
)
from backend.utils.vision import VisionImage, VisionUnsupportedError
from backend.worker.tasks.documents import _transcribe_batch


def _png(color=(40, 120, 200), size=(32, 32)) -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def _page(number: int, *, with_image: bool = True) -> PageSource:
    # A real PNG: bytes Pillow cannot open are left out, never sent (#714).
    return PageSource(
        page_number=number,
        text=f"structural text {number}",
        images=(
            [VisionImage(data=_png(), media_type="image/png")] if with_image else []
        ),
    )


class _Backend:
    """Records calls and replays a scripted outcome per page."""

    def __init__(self, outcomes=None):
        self.outcomes = outcomes or {}
        self.calls = []

    def generate_text_from_images(self, prompt, images, timeout=120, max_tokens=8192):
        self.calls.append(prompt)
        # The page number is recoverable from the structural text in the prompt.
        for number, outcome in self.outcomes.items():
            if f"structural text {number}" in prompt:
                if isinstance(outcome, Exception):
                    raise outcome
                return outcome
        return "transcribed"


# ---------------------------------------------------------------------------
# Fan-out
# ---------------------------------------------------------------------------


def test_every_page_with_an_image_is_transcribed():
    backend = _Backend()
    results = _transcribe_batch(
        backend, [_page(1), _page(2), _page(3)], is_rendered=True
    )
    assert results == {1: "transcribed", 2: "transcribed", 3: "transcribed"}
    assert len(backend.calls) == 3


def test_pages_without_images_are_not_sent():
    backend = _Backend()
    results = _transcribe_batch(backend, [_page(1, with_image=False)], is_rendered=True)
    assert results == {}
    assert backend.calls == []


def test_an_empty_batch_makes_no_calls():
    backend = _Backend()
    assert _transcribe_batch(backend, [], is_rendered=True) == {}
    assert backend.calls == []


# ---------------------------------------------------------------------------
# Failure handling
# ---------------------------------------------------------------------------


def test_one_page_failing_does_not_lose_the_others():
    """A transient provider error on page 2 must not abandon a long document;
    that page falls back to its structural text."""
    backend = _Backend({2: RuntimeError("provider hiccup")})
    results = _transcribe_batch(
        backend, [_page(1), _page(2), _page(3)], is_rendered=True
    )
    assert set(results) == {1, 3}


def test_a_page_returning_nothing_is_omitted_rather_than_stored_blank():
    backend = _Backend({2: "   "})
    results = _transcribe_batch(backend, [_page(1), _page(2)], is_rendered=True)
    assert set(results) == {1}


def test_vision_unsupported_propagates_so_the_document_downgrades_once():
    """Unlike a per-page failure, this condemns every remaining page. It must
    reach the caller so the whole document is downgraded and warned about,
    rather than retried uselessly on every page in turn."""
    backend = _Backend({1: VisionUnsupportedError("no vision on this model")})
    with pytest.raises(VisionUnsupportedError):
        _transcribe_batch(backend, [_page(1), _page(2)], is_rendered=True)


# ---------------------------------------------------------------------------
# Capability detection
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "message",
    [
        "400 does not support image input",
        "This model does not support vision",
        "Invalid content type: image_url is not supported",
        "multimodal input rejected",
    ],
)
def test_provider_refusals_are_recognised_as_missing_vision(message):
    assert is_vision_unsupported_error(RuntimeError(message))


@pytest.mark.parametrize(
    "message",
    [
        "429 rate limit exceeded",
        "connection reset by peer",
        "500 internal server error",
    ],
)
def test_transient_errors_are_not_mistaken_for_missing_vision(message):
    """Getting this wrong would permanently downgrade a document over one
    flaky call."""
    assert not is_vision_unsupported_error(RuntimeError(message))


def test_the_default_backend_reports_no_vision_and_no_capability_answer():
    backend = LLMBackend()
    assert backend.supports_vision() is None
    with pytest.raises(VisionUnsupportedError):
        backend.generate_text_from_images("prompt", [])


# ---------------------------------------------------------------------------
# Stalled-parse recovery
# ---------------------------------------------------------------------------


def _document(status, minutes_since_write):
    from datetime import timedelta

    from backend.models.document import (
        Document,
        DocumentStatus,
        parse_looks_stalled,
    )
    from backend.utils.time import utc_now

    document = Document(
        recording_id=1,
        title="deck.pdf",
        file_path="/data/documents/deck.pdf",
        status=getattr(DocumentStatus, status),
    )
    document.updated_at = utc_now() - timedelta(minutes=minutes_since_write)
    return parse_looks_stalled(document)


def test_a_running_parse_is_not_treated_as_stalled():
    assert _document("PROCESSING", 1) is False


def test_a_long_silent_parse_is_treated_as_stalled():
    """A worker that dies -- OOM, restart, segfault -- leaves the row PROCESSING
    forever. Without this the document is unrecoverable from the UI, which is a
    worse failure than the duplicate parse the guard prevents."""
    assert _document("PROCESSING", 30) is True


def test_a_finished_document_is_never_stalled():
    for status in ("READY", "ERROR", "PENDING"):
        assert _document(status, 999) is False


# ---------------------------------------------------------------------------
# Backend resolution
# ---------------------------------------------------------------------------


def test_the_llm_factory_is_importable_from_the_task_module():
    """Regression: the factory was reached through `from .constants import *`,
    but constants imports it inside a function, so it never entered the star
    namespace. Every document silently downgraded to a structural parse with a
    warning blaming the user's provider settings."""
    import backend.worker.tasks.documents as documents_task

    source = documents_task._resolve_backend_for_document.__code__
    referenced = set(source.co_names)
    assert "get_llm_backend_with_secondary" in referenced

    # The name must resolve for real, not just appear in the source.
    from backend.processing.llm_services import (  # noqa: F401
        get_llm_backend_with_secondary,
    )


def test_no_backend_reports_a_reason_the_user_can_act_on(monkeypatch):
    """The reason is rendered on the document card, so an internal fault must
    not read as "no AI provider is configured" and send the user to settings
    that were never wrong."""
    import backend.worker.tasks.documents as documents_task

    class _Session:
        def get(self, model, pk):
            return None

    class _Doc:
        recording_id = 1

    backend, reason = documents_task._resolve_backend_for_document(_Session(), _Doc())
    assert backend is None
    assert reason and reason.endswith(".")
    assert "not linked to a user account" in reason


# ---------------------------------------------------------------------------
# Provider-chain forwarding
# ---------------------------------------------------------------------------


class _VisionBackend(LLMBackend):
    def __init__(self, answer, result="described"):
        self.answer = answer
        self.result = result
        self.calls = 0

    def supports_vision(self):
        return self.answer

    def generate_text_from_images(self, prompt, images, timeout=120, max_tokens=8192):
        self.calls += 1
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


def _chain(primary, secondary):
    from backend.processing.llm_backends.factory import SecondaryLLMBackend

    return SecondaryLLMBackend(primary, secondary)


def test_the_provider_chain_forwards_image_calls_to_the_primary():
    """Regression: SecondaryLLMBackend did not override this, so the call fell
    through to LLMBackend's default, raised VisionUnsupportedError, and every
    document downgraded even when the primary handled images fine."""
    primary = _VisionBackend(None, "from primary")
    chain = _chain(primary, _VisionBackend(None, "from secondary"))

    assert chain.generate_text_from_images("prompt", []) == "from primary"
    assert primary.calls == 1


def test_a_primary_without_vision_falls_through_to_the_secondary():
    primary = _VisionBackend(None, VisionUnsupportedError("no vision"))
    secondary = _VisionBackend(None, "from secondary")
    chain = _chain(primary, secondary)

    assert chain.generate_text_from_images("prompt", []) == "from secondary"
    assert secondary.calls == 1


def test_both_failing_surfaces_the_unsupported_error_so_the_document_downgrades():
    chain = _chain(
        _VisionBackend(None, VisionUnsupportedError("primary has no vision")),
        _VisionBackend(None, VisionUnsupportedError("secondary has no vision")),
    )
    with pytest.raises(VisionUnsupportedError):
        chain.generate_text_from_images("prompt", [])


@pytest.mark.parametrize(
    "primary,secondary,expected",
    [
        (True, False, True),
        (False, True, True),
        (False, False, False),
        (None, False, None),
        (None, None, None),
    ],
)
def test_chain_vision_capability_is_tri_state(primary, secondary, expected):
    chain = _chain(_VisionBackend(primary), _VisionBackend(secondary))
    assert chain.supports_vision() is expected


# ---------------------------------------------------------------------------
# Stall reporting crosses to the client as a boolean
# ---------------------------------------------------------------------------


def test_serialized_documents_carry_a_server_computed_stall_flag():
    """Regression: the client derived this from updated_at, which serialises
    without a timezone. JavaScript parses such a string as local time, so any
    browser outside UTC misjudged the age by its whole offset and showed every
    running parse as stalled while the server refused the re-parse it invited."""
    from datetime import timedelta

    from backend.models.document import Document, DocumentStatus
    from backend.models.recording_public import serialize_document
    from backend.utils.time import utc_now

    document = Document(
        recording_id=1,
        title="deck.pdf",
        file_path="/data/documents/deck.pdf",
        status=DocumentStatus.PROCESSING,
    )
    document.id = 1

    document.updated_at = utc_now() - timedelta(seconds=30)
    assert serialize_document(document, recording_public_id="x").is_stalled is False

    document.updated_at = utc_now() - timedelta(minutes=45)
    assert serialize_document(document, recording_public_id="x").is_stalled is True


def test_the_endpoint_and_the_serializer_share_one_stall_definition():
    """They gated on separate copies before, so the button could refuse a
    re-parse the label was inviting."""
    from backend.api.v1.endpoints import documents as documents_api
    from backend.models.document import parse_looks_stalled

    assert documents_api.parse_looks_stalled is parse_looks_stalled


# ---------------------------------------------------------------------------
# Progress granularity
# ---------------------------------------------------------------------------


class _RecordingSession:
    """Captures pages_parsed at each commit, which is what the UI polls."""

    def __init__(self):
        self.snapshots: list[int] = []
        self._pending = None

    def add(self, obj):
        if hasattr(obj, "pages_parsed"):
            self._pending = obj

    def commit(self):
        if self._pending is not None:
            self.snapshots.append(self._pending.pages_parsed)


def _progress_writer(monkeypatch, *, use_vision, transcriptions=None, ocr=None):
    import backend.worker.tasks.documents as documents_task

    monkeypatch.setattr(documents_task, "ocr_images", lambda _images: ocr)
    if transcriptions is not None:

        def _fake_batch(
            backend, pages, *, is_rendered, on_page_done=None, **_callbacks
        ):
            # Mirror the real contract: settle one page at a time, by number.
            for page in pages:
                if on_page_done is not None:
                    on_page_done(page.page_number)
            return transcriptions

        monkeypatch.setattr(documents_task, "_transcribe_batch", _fake_batch)

    class _Doc:
        id = 1
        pages_parsed = 0

    session = _RecordingSession()
    writer = documents_task._PageWriter(
        session,
        _Doc(),
        backend=object(),
        use_vision=use_vision,
        is_rendered=True,
        already_parsed=0,
    )
    return writer, session


def test_progress_is_reported_per_page_not_per_batch(monkeypatch):
    """The whole batch used to persist in one write, so a seven-page document
    went straight from 0 to 7 with nothing in between."""
    pages = [_page(n) for n in range(1, 8)]
    writer, session = _progress_writer(
        monkeypatch,
        use_vision=True,
        transcriptions={n: f"page {n}" for n in range(1, 8)},
    )

    writer.flush(pages)

    # One report per page, in order, and no redundant repeat at the end.
    assert session.snapshots == [1, 2, 3, 4, 5, 6, 7]
    assert writer._document.pages_parsed == 7


def test_progress_never_goes_backwards_or_exceeds_the_batch(monkeypatch):
    pages = [_page(n) for n in range(1, 6)]
    writer, session = _progress_writer(
        monkeypatch,
        use_vision=True,
        transcriptions={n: f"page {n}" for n in range(1, 6)},
    )

    writer.flush(pages)

    assert session.snapshots == sorted(session.snapshots)
    assert max(session.snapshots) == len(pages)


def test_pages_that_skip_the_vision_tier_still_report_progress(monkeypatch):
    """Structural and OCR pages never reach the vision callback, so without
    their own report a text-only document would show no movement at all."""
    pages = [_page(n, with_image=False) for n in range(1, 5)]
    writer, session = _progress_writer(monkeypatch, use_vision=False)

    writer.flush(pages)

    assert session.snapshots == [1, 2, 3, 4]


def test_a_page_is_never_counted_twice(monkeypatch):
    """One page with an image and one without: the first reports through the
    vision callback, the second at persist. Counting both paths for the same
    page would overshoot the page count and break the percentage."""
    pages = [_page(1), _page(2, with_image=False)]
    writer, session = _progress_writer(
        monkeypatch, use_vision=True, transcriptions={1: "page 1"}
    )

    writer.flush(pages)

    assert max(session.snapshots) == 2
    assert writer._document.pages_parsed == 2


def test_progress_accumulates_across_batches(monkeypatch):
    writer, session = _progress_writer(
        monkeypatch, use_vision=True, transcriptions={1: "a", 2: "b"}
    )

    writer.flush([_page(1), _page(2)])
    first = writer._document.pages_parsed
    writer.flush([_page(1), _page(2)])

    assert first == 2
    assert writer._document.pages_parsed == 4


# ---------------------------------------------------------------------------
# Image size limits (#698)
# ---------------------------------------------------------------------------


class _ClaudeLikeBackend:
    """Refuses an image as the Claude API does: past 5 MB of base64, or past
    8000 px a side. Records the size of every image it accepts."""

    BASE64_LIMIT = 5 * 1024 * 1024

    def __init__(self, result="A photo of a whiteboard."):
        self.result = result
        self.seen: list[tuple[int, int, int, str]] = []

    def generate_text_from_images(self, prompt, images, timeout=120, max_tokens=8192):
        import io

        from PIL import Image

        for image in images:
            with Image.open(io.BytesIO(image.data)) as picture:
                width, height = picture.size
            if len(image.to_base64()) > self.BASE64_LIMIT or max(width, height) > 8000:
                raise RuntimeError(
                    "400 invalid_request_error: image exceeds 5 MB maximum"
                )
            self.seen.append((width, height, len(image.data), image.media_type))
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


def _large_photo(path) -> bytes:
    """An 11 MB PNG of noise, 2700 x 1400 px: past every provider's limit."""
    import io
    import os

    from PIL import Image

    buffer = io.BytesIO()
    Image.frombytes("RGB", (2700, 1400), os.urandom(2700 * 1400 * 3)).save(
        buffer, format="PNG", compress_level=1
    )
    path.write_bytes(buffer.getvalue())
    return buffer.getvalue()


@pytest.fixture
def image_parse(monkeypatch, tmp_path):
    """Parses one PNG through the real orchestrator on SQLite, with a stub
    model and stub OCR. Returns (run, session); run(backend, ocr) parses."""
    from sqlalchemy import text
    from sqlmodel import Session, create_engine

    import backend.worker.tasks.documents as documents_task
    from backend.models.document import Document, DocumentStatus
    from backend.tests.sqlite_schemas import DOCUMENT_PAGES_SCHEMA, DOCUMENTS_SCHEMA

    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        for schema in (DOCUMENTS_SCHEMA, DOCUMENT_PAGES_SCHEMA):
            connection.execute(text(schema))
    session = Session(engine)
    path = tmp_path / "whiteboard.png"
    original = _large_photo(path)
    document = Document(
        id=1,
        recording_id=10,
        title="whiteboard.png",
        file_path=str(path),
        file_type="image/png",
        status=DocumentStatus.PROCESSING,
    )
    session.add(document)
    session.commit()
    monkeypatch.setattr(documents_task, "_index_page", lambda *_args: 0)
    monkeypatch.setattr(documents_task, "_flag_notes_stale", lambda *_args: None)

    def run(backend, *, ocr=None, ocr_installed=True):
        monkeypatch.setattr(
            documents_task, "_resolve_vision", lambda *_args: (backend, True, None)
        )
        monkeypatch.setattr(documents_task, "ocr_images", lambda _images: ocr)
        monkeypatch.setattr(documents_task, "ocr_is_available", lambda: ocr_installed)
        documents_task._parse_document(session, document, str(path))
        assert path.read_bytes() == original, "the stored file must never change"
        return document

    yield run, session
    session.close()


def test_a_large_image_is_read_by_the_model_not_dropped_to_ocr(image_parse):
    """Regression for #698: the upload went to the model as it was, the Claude
    API refused it, and the page fell back to OCR with no word on the card."""
    from backend.models.document import DocumentStatus
    from backend.models.document_page import DocumentPage

    run, session = image_parse
    backend = _ClaudeLikeBackend()

    document = run(backend, ocr="OCR text that must not be used")

    assert document.status == DocumentStatus.READY
    assert document.parse_warning is None
    page = session.get(DocumentPage, 1)
    assert (page.parse_mode, page.content) == (
        PageParseMode.VISUAL,
        "A photo of a whiteboard.",
    )
    ((width, height, size_bytes, _media_type),) = backend.seen
    assert max(width, height) <= 2576
    assert size_bytes <= 3_500_000


def test_a_failed_vision_call_is_reported_on_the_card(image_parse):
    from backend.models.document import DocumentStatus

    run, _session = image_parse

    document = run(
        _ClaudeLikeBackend(result=RuntimeError("503 overloaded")),
        ocr="Text recovered by local OCR from the whiteboard.",
    )

    assert document.status == DocumentStatus.READY
    assert document.parse_warning.startswith("Visual analysis failed, so this file")
    assert "local OCR" in document.parse_warning


def test_an_image_whose_vision_call_failed_says_so_in_its_error(image_parse):
    """It used to say no vision-capable model was available, which sends the
    user to settings that were never wrong."""
    run, _session = image_parse

    with pytest.raises(RuntimeError) as raised:
        run(
            _ClaudeLikeBackend(result=RuntimeError("503 overloaded")),
            ocr_installed=False,
        )

    assert "visual analysis by the AI model failed" in str(raised.value)
    assert "no vision-capable" not in str(raised.value)


def test_a_document_counts_the_pages_whose_vision_call_failed(monkeypatch):
    import backend.worker.tasks.documents as documents_task

    def _batch(backend, pages, *, is_rendered, on_page_failed=None, **_callbacks):
        for page in pages:
            if page.page_number in (2, 5):
                on_page_failed(page.page_number)
        return {page.page_number: "described" for page in pages}

    writer, _session = _progress_writer(monkeypatch, use_vision=True)
    monkeypatch.setattr(documents_task, "_transcribe_batch", _batch)
    writer.flush([_page(n) for n in range(1, 7)])

    assert writer.vision_failures == 2
    warning = documents_task._parse_warning(None, writer, 9)
    assert warning.startswith("Visual analysis failed on 2 of 9 pages, so those pages")
    writer.vision_failures = 1
    assert "1 of 9 pages, so that page holds" in documents_task._parse_warning(
        None, writer, 9
    )


def test_a_failing_page_is_reported_to_the_caller():
    failed: list[int] = []
    backend = _Backend({2: RuntimeError("provider hiccup")})

    _transcribe_batch(
        backend,
        [_page(1), _page(2), _page(3)],
        is_rendered=True,
        on_page_failed=failed.append,
    )

    assert failed == [2]


def test_the_model_gets_a_fitted_copy_and_the_page_keeps_its_original(tmp_path):
    from backend.processing.documents.vision import transcribe_page

    original = _large_photo(tmp_path / "photo.png")
    page = PageSource(
        page_number=1,
        images=[VisionImage(data=original, media_type="image/png")],
    )
    backend = _ClaudeLikeBackend()

    assert transcribe_page(backend, page, is_rendered_page=True).text
    assert page.images[0].data == original
    assert backend.seen[0][3] == "image/jpeg"


# ---------------------------------------------------------------------------
# A page's pictures in several calls (#714)
# ---------------------------------------------------------------------------


class _ClaudeRequestBackend:
    """Refuses a call as the Claude API does: past 32 MB a request, past 5 MB
    of base64 an image, or carrying an image it cannot decode. Records the
    labels of the pictures in each call it answers, and fails the calls
    numbered in ``fail_calls`` as an overloaded provider would."""

    REQUEST_LIMIT = 32 * 1000 * 1000
    IMAGE_LIMIT = 5 * 1024 * 1024

    def __init__(self, fail_calls=()):
        self.fail_calls = set(fail_calls)
        self.attempts = 0
        self.calls: list[list[str]] = []
        self.request_sizes: list[int] = []

    def generate_text_from_images(self, prompt, images, timeout=120, max_tokens=8192):
        from backend.tests.picture_fixtures import label_of

        self.attempts += 1
        size = len(prompt) + sum(len(image.to_base64()) for image in images)
        if size > self.REQUEST_LIMIT:
            raise RuntimeError("413 request_too_large: request exceeds 32 MB")
        labels = []
        for image in images:
            if len(image.to_base64()) > self.IMAGE_LIMIT:
                raise RuntimeError("400 invalid_request_error: image exceeds 5 MB")
            try:
                labels.append(label_of(image.data))
            except OSError as exc:
                raise RuntimeError(
                    "400 invalid_request_error: Could not process image"
                ) from exc
        if self.attempts in self.fail_calls:
            raise RuntimeError("529 overloaded_error")
        self.calls.append(labels)
        self.request_sizes.append(size)
        return f"Pictures {', '.join(labels)}."


def _pictures(count: int) -> list[VisionImage]:
    """``count`` small JPEGs labelled 1, 2, 3 and on."""
    from backend.tests.picture_fixtures import jpeg, labelled

    base = jpeg()
    return [
        VisionImage(data=labelled(base, str(number)), media_type="image/jpeg")
        for number in range(1, count + 1)
    ]


def _described(first: int, last: int) -> str:
    """The stub's answer for a call carrying pictures ``first`` to ``last``."""
    return f"Pictures {', '.join(str(n) for n in range(first, last + 1))}."


def _emf() -> VisionImage:
    from backend.tests.picture_fixtures import emf_drawing

    return VisionImage(data=emf_drawing(), media_type="image/x-emf")


def _figures(images) -> PageSource:
    return PageSource(page_number=1, title="Site visit", text="", images=images)


def _transcribe(backend, images):
    from backend.processing.documents.vision import transcribe_page

    return transcribe_page(backend, _figures(images), is_rendered_page=False)


def test_a_page_of_many_pictures_is_read_in_calls_joined_in_order():
    from backend.processing.documents.vision_input import MAX_BATCH_IMAGES

    count = MAX_BATCH_IMAGES * 2 + 3
    backend = _ClaudeRequestBackend()

    transcription = _transcribe(backend, _pictures(count))

    numbers = [str(n) for n in range(1, count + 1)]
    assert backend.calls == [
        numbers[:MAX_BATCH_IMAGES],
        numbers[MAX_BATCH_IMAGES : 2 * MAX_BATCH_IMAGES],
        numbers[2 * MAX_BATCH_IMAGES :],
    ]
    assert transcription.text == "\n\n".join(
        f"Pictures {', '.join(call)}." for call in backend.calls
    )
    assert (transcription.left_out, transcription.failed) == (0, 0)


def test_a_failed_call_costs_only_its_own_pictures():
    from backend.processing.documents.vision_input import MAX_BATCH_IMAGES as per_call

    backend = _ClaudeRequestBackend(fail_calls={2})

    transcription = _transcribe(backend, _pictures(per_call * 2 + 3))

    assert transcription.text == (
        f"{_described(1, per_call)}\n\n{_described(per_call * 2 + 1, per_call * 2 + 3)}"
    )
    assert transcription.failed == per_call
    assert not transcription.page_failed


def test_a_page_fails_only_when_every_call_did():
    backend = _ClaudeRequestBackend(fail_calls={1, 2})

    transcription = _transcribe(backend, _pictures(10))

    assert (transcription.text, transcription.failed) == (None, 10)
    assert transcription.page_failed


def test_a_picture_pillow_cannot_read_is_left_out_not_sent():
    """An EMF drawing went with its neighbours, the provider refused it, and
    the whole call failed (#714)."""
    first, second = _pictures(2)
    backend = _ClaudeRequestBackend()

    transcription = _transcribe(backend, [first, _emf(), second])

    assert backend.calls == [["1", "2"]]
    assert transcription.text == _described(1, 2)
    assert transcription.left_out == 1


def test_a_page_of_only_unreadable_pictures_makes_no_call():
    backend = _ClaudeRequestBackend()

    transcription = _transcribe(backend, [_emf(), _emf()])

    assert backend.attempts == 0
    assert (transcription.text, transcription.left_out) == (None, 2)
    assert not transcription.page_failed


class _FailsSomeCalls(_ClaudeRequestBackend):
    """Fails page 3's call and page 2's second one. Pages run concurrently, so
    a call is picked out by its page and its first picture, not by its order."""

    def __init__(self, second_call_starts_at: int):
        super().__init__()
        self.second_call_starts_at = str(second_call_starts_at)

    def generate_text_from_images(self, prompt, images, timeout=120, max_tokens=8192):
        from backend.tests.picture_fixtures import label_of

        first = label_of(images[0].data)
        if "page 3" in prompt or (
            "page 2" in prompt and first == self.second_call_starts_at
        ):
            raise RuntimeError("529 overloaded_error")
        return super().generate_text_from_images(prompt, images, timeout, max_tokens)


def test_the_fan_out_reports_what_each_page_did_not_describe():
    from backend.processing.documents.vision_input import MAX_BATCH_IMAGES as per_call

    pages = [
        _figures([*_pictures(1), _emf()]),
        PageSource(page_number=2, images=_pictures(per_call + 2)),
        PageSource(page_number=3, images=_pictures(1)),
    ]
    failed_pages: list[int] = []
    missed: list[tuple[int, int]] = []

    results = _transcribe_batch(
        _FailsSomeCalls(second_call_starts_at=per_call + 1),
        pages,
        is_rendered=False,
        on_page_failed=failed_pages.append,
        on_pictures_missed=lambda left_out, failed: missed.append((left_out, failed)),
    )

    assert results == {1: _described(1, 1), 2: _described(1, per_call)}
    assert failed_pages == [3]
    # (left out, failed) per page: page 3's picture counts with its failed page.
    assert sorted(missed) == [(0, 0), (0, 2), (1, 0)]


def _writer_with(monkeypatch, *, pictures_failed=0, left_out=0, vision_failures=0):
    writer, _session = _progress_writer(monkeypatch, use_vision=True)
    writer.pictures_failed = pictures_failed
    writer.pictures_left_out = left_out
    writer.vision_failures = vision_failures
    return writer


def test_the_warning_counts_pictures_failed_and_left_out(monkeypatch):
    import backend.worker.tasks.documents as documents_task

    writer = _writer_with(monkeypatch, pictures_failed=8, left_out=2)

    assert documents_task._parse_warning(None, writer, 3) == (
        "Visual analysis failed on 8 pictures, so they are not described. Check "
        "the server logs for the reason. Visual analysis left out 2 pictures "
        "Nojoin cannot open, such as EMF or WMF drawings."
    )
    one = _writer_with(monkeypatch, pictures_failed=1, left_out=1)
    assert documents_task._parse_warning(None, one, 3) == (
        "Visual analysis failed on 1 picture, so it is not described. Check the "
        "server logs for the reason. Visual analysis left out 1 picture Nojoin "
        "cannot open, such as an EMF or WMF drawing."
    )


def test_the_warning_names_the_server_logs_once(monkeypatch):
    import backend.worker.tasks.documents as documents_task

    writer = _writer_with(monkeypatch, pictures_failed=3, vision_failures=1)

    warning = documents_task._parse_warning(None, writer, 9)

    assert warning.startswith("Visual analysis failed on 1 of 9 pages, so that page")
    assert warning.endswith(
        "Visual analysis failed on 3 pictures, so they are not described."
    )
    assert warning.count("Check the server logs") == 1


def test_a_clean_parse_has_no_warning(monkeypatch):
    import backend.worker.tasks.documents as documents_task

    assert documents_task._parse_warning(None, _writer_with(monkeypatch), 3) is None


@pytest.fixture
def parse_file(monkeypatch):
    """Parses a file through the real orchestrator on SQLite, with a stub model
    and no OCR. run(path, backend) returns the document and its pages."""
    from sqlalchemy import text
    from sqlmodel import Session, create_engine, select

    import backend.worker.tasks.documents as documents_task
    from backend.models.document import Document, DocumentStatus
    from backend.models.document_page import DocumentPage
    from backend.tests.sqlite_schemas import DOCUMENT_PAGES_SCHEMA, DOCUMENTS_SCHEMA

    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        for schema in (DOCUMENTS_SCHEMA, DOCUMENT_PAGES_SCHEMA):
            connection.execute(text(schema))
    session = Session(engine)
    monkeypatch.setattr(documents_task, "_index_page", lambda *_args: 0)
    monkeypatch.setattr(documents_task, "_flag_notes_stale", lambda *_args: None)
    monkeypatch.setattr(documents_task, "ocr_images", lambda _images: None)
    monkeypatch.setattr(documents_task, "ocr_is_available", lambda: False)

    def run(path, backend):
        monkeypatch.setattr(
            documents_task, "_resolve_vision", lambda *_args: (backend, True, None)
        )
        document = Document(
            id=1,
            recording_id=10,
            title=path.name,
            file_path=str(path),
            file_type="application/octet-stream",
            status=DocumentStatus.PROCESSING,
        )
        session.add(document)
        session.commit()
        documents_task._parse_document(session, document, str(path))
        pages = session.exec(
            select(DocumentPage).order_by(DocumentPage.page_number)
        ).all()
        return document, pages

    yield run
    session.close()


def _word_file(path, photos: list[bytes], *, extra_parts=()):
    """A Word file: a heading, a line of text, then ``photos`` in order, with
    ``extra_parts`` ((before_photo, file_name, content_type, blob)) among them."""
    docx = pytest.importorskip("docx")
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    document = docx.Document()
    document.add_heading("Site visit", level=1)
    document.add_paragraph("Photos from the walk round the east wing.")
    for index, photo in enumerate(photos):
        for before, file_name, content_type, blob in extra_parts:
            if before == index:
                name = PackURI(f"/word/media/{file_name}")
                part = Part(name, content_type, blob, document.part.package)
                document.part.relate_to(part, RELATIONSHIP_TYPE.IMAGE)
        document.add_picture(io.BytesIO(photo))
    document.save(str(path))


def test_a_word_file_full_of_photos_gets_every_description(parse_file, tmp_path):
    """Regression for #714: nine photos of about 3 MB went to the model in one
    call, past the Claude API's 32 MB request limit, and the section lost all
    of its visual analysis. An EMF drawing beside them failed the call too."""
    from backend.models.document import DocumentStatus
    from backend.processing.documents.vision_input import MAX_IMAGE_BYTES
    from backend.tests.picture_fixtures import emf_drawing, jpeg, labelled

    base = jpeg((2200, 1500), noise=True)
    assert 2_500_000 < len(base) < MAX_IMAGE_BYTES  # sent as it is
    photos = [labelled(base, str(number)) for number in range(1, 10)]
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="8" height="8"/>'
    path = tmp_path / "site-visit.docx"
    _word_file(
        path,
        photos,
        extra_parts=[
            (5, "drawing.emf", "image/x-emf", emf_drawing()),
            (7, "icon.svg", "image/svg+xml", svg),
        ],
    )
    backend = _ClaudeRequestBackend()

    document, (page,) = parse_file(path, backend)

    assert document.status == DocumentStatus.READY
    assert page.parse_mode == PageParseMode.VISUAL
    # Every photo, in order, in as few calls as fit the tightest request
    # limit: Gemini's 20 MB of inline data.
    assert [label for call in backend.calls for label in call] == [
        str(number) for number in range(1, 10)
    ]
    from backend.processing.documents.vision_input import (
        MAX_BATCH_BYTES,
        MAX_BATCH_IMAGES,
    )

    per_call = min(MAX_BATCH_BYTES // len(photos[0]), MAX_BATCH_IMAGES)
    assert len(backend.calls) == -(-len(photos) // per_call)
    assert max(backend.request_sizes) <= 20_000_000
    described = "\n\n".join(f"Pictures {', '.join(call)}." for call in backend.calls)
    assert page.content == f"Photos from the walk round the east wing.\n\n{described}"
    # The EMF is counted. The SVG is neither sent nor counted: Word stores a
    # PNG of every SVG picture as well, and that copy is the one described.
    assert document.parse_warning == (
        "Visual analysis left out 1 picture Nojoin cannot open, such as an EMF "
        "or WMF drawing."
    )


def test_an_image_pillow_cannot_open_says_so_in_its_error(parse_file, tmp_path):
    path = tmp_path / "scan.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n and then nothing")
    backend = _ClaudeRequestBackend()

    with pytest.raises(RuntimeError) as raised:
        parse_file(path, backend)

    assert str(raised.value) == (
        "This image could not be read: the file is not a picture Nojoin can open."
    )
    assert backend.attempts == 0
